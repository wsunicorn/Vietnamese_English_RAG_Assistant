from pathlib import Path
from zipfile import ZipFile

from app.ingestion.chunking import DocumentPage, ParsedDocument, normalize_whitespace

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".zip"}


class UnsupportedDocumentError(ValueError):
    pass


class DocumentParser:
    def parse_many(self, path: Path, *, filename: str, content_type: str) -> list[ParsedDocument]:
        extension = path.suffix.lower()
        if extension == ".zip":
            return self._parse_notion_zip(path, filename=filename, content_type=content_type)
        return [self.parse(path, filename=filename, content_type=content_type)]

    def parse(self, path: Path, *, filename: str, content_type: str) -> ParsedDocument:
        extension = path.suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise UnsupportedDocumentError(f"Unsupported file extension: {extension}")

        if extension == ".txt":
            return self._parse_txt(path, filename=filename, content_type=content_type)
        if extension in {".md", ".markdown"}:
            return self._parse_markdown(path, filename=filename, content_type=content_type)
        if extension == ".pdf":
            return self._parse_pdf(path, filename=filename, content_type=content_type)
        if extension == ".docx":
            return self._parse_docx(path, filename=filename, content_type=content_type)
        raise UnsupportedDocumentError(f"Unsupported file extension: {extension}")

    def _parse_txt(self, path: Path, *, filename: str, content_type: str) -> ParsedDocument:
        text = repair_mojibake(path.read_text(encoding="utf-8", errors="ignore"))
        return ParsedDocument(
            filename=filename,
            pages=[DocumentPage(page_number=1, text=text, metadata={"parser": "plain-text"})],
            metadata={"content_type": content_type, "source_type": "txt"},
            language=detect_language(text),
        )

    def _parse_markdown(self, path: Path, *, filename: str, content_type: str) -> ParsedDocument:
        text = repair_mojibake(path.read_text(encoding="utf-8", errors="ignore"))
        return parsed_markdown_document(
            text,
            filename=filename,
            content_type=content_type,
            source_type="markdown",
            source_path=filename,
        )

    def _parse_pdf(self, path: Path, *, filename: str, content_type: str) -> ParsedDocument:
        pages: list[DocumentPage] = []
        parser = "pypdf"
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            for index, page in enumerate(reader.pages, start=1):
                text = repair_mojibake(page.extract_text() or "")
                if normalize_whitespace(text):
                    pages.append(
                        DocumentPage(
                            page_number=index,
                            text=text,
                            metadata={"parser": parser, "section": f"Page {index}"},
                        )
                    )
        except Exception:
            pages = []

        if not pages:
            text = repair_mojibake(self._parse_with_docling(path))
            pages = [DocumentPage(page_number=1, text=text, metadata={"parser": "docling"})]
            parser = "docling"

        full_text = "\n".join(page.text for page in pages)
        return ParsedDocument(
            filename=filename,
            pages=pages,
            metadata={"content_type": content_type, "source_type": "pdf", "parser": parser},
            language=detect_language(full_text),
        )

    def _parse_docx(self, path: Path, *, filename: str, content_type: str) -> ParsedDocument:
        parser = "python-docx"
        try:
            from docx import Document

            doc = Document(str(path))
            paragraphs = [paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()]
            text = repair_mojibake("\n\n".join(paragraphs))
        except Exception:
            text = repair_mojibake(self._parse_with_docling(path))
            parser = "docling"

        return ParsedDocument(
            filename=filename,
            pages=[
                DocumentPage(
                    page_number=1,
                    text=text,
                    metadata={"parser": parser, "section": "Document body"},
                )
            ],
            metadata={"content_type": content_type, "source_type": "docx", "parser": parser},
            language=detect_language(text),
        )

    @staticmethod
    def _parse_with_docling(path: Path) -> str:
        try:
            from docling.document_converter import DocumentConverter

            converter = DocumentConverter()
            result = converter.convert(str(path))
            return result.document.export_to_markdown()
        except Exception as exc:
            raise RuntimeError(f"Docling could not parse document: {exc}") from exc

    def _parse_notion_zip(
        self, path: Path, *, filename: str, content_type: str
    ) -> list[ParsedDocument]:
        documents: list[ParsedDocument] = []
        with ZipFile(path) as archive:
            for info in archive.infolist():
                if info.is_dir() or should_ignore_zip_member(info.filename):
                    continue
                extension = Path(info.filename).suffix.lower()
                if extension not in {".md", ".markdown", ".html", ".htm", ".txt"}:
                    continue
                raw = archive.read(info)
                text = repair_mojibake(raw.decode("utf-8", errors="ignore"))
                if extension in {".html", ".htm"}:
                    text = html_to_text(text)
                title = source_title_from_path(info.filename)
                if extension in {".md", ".markdown"}:
                    parsed = parsed_markdown_document(
                        text,
                        filename=title,
                        content_type="text/markdown",
                        source_type="notion",
                        source_path=info.filename,
                    )
                else:
                    parsed = ParsedDocument(
                        filename=title,
                        pages=[
                            DocumentPage(
                                page_number=None,
                                text=text,
                                metadata={
                                    "parser": "notion-export",
                                    "section": title,
                                    "source_path": info.filename,
                                },
                            )
                        ],
                        metadata={
                            "content_type": content_type,
                            "source_type": "notion",
                            "source_path": info.filename,
                            "source_title": title,
                            "notion_export": filename,
                        },
                        language=detect_language(text),
                    )
                parsed.metadata["notion_export"] = filename
                documents.append(parsed)

        if not documents:
            raise UnsupportedDocumentError("Notion export ZIP did not contain Markdown, HTML, or TXT pages.")
        return documents


def parsed_markdown_document(
    text: str,
    *,
    filename: str,
    content_type: str,
    source_type: str,
    source_path: str,
) -> ParsedDocument:
    pages = markdown_sections_to_pages(text)
    full_text = "\n".join(page.text for page in pages)
    return ParsedDocument(
        filename=filename,
        pages=pages,
        metadata={
            "content_type": content_type,
            "source_type": source_type,
            "source_path": source_path,
            "source_title": filename,
            "parser": "markdown",
        },
        language=detect_language(full_text),
    )


def markdown_sections_to_pages(text: str) -> list[DocumentPage]:
    pages: list[DocumentPage] = []
    current_title = "Document body"
    current_lines: list[str] = []

    def flush() -> None:
        body = normalize_whitespace("\n".join(current_lines))
        if body:
            pages.append(
                DocumentPage(
                    page_number=None,
                    text=body,
                    metadata={"parser": "markdown", "section": current_title},
                )
            )

    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("#"):
            heading = stripped.lstrip("#").strip()
            if heading:
                flush()
                current_title = heading[:255]
                current_lines = [heading]
                continue
        current_lines.append(line)

    flush()
    if not pages and normalize_whitespace(text):
        pages.append(
            DocumentPage(
                page_number=None,
                text=text,
                metadata={"parser": "markdown", "section": current_title},
            )
        )
    return pages


def should_ignore_zip_member(name: str) -> bool:
    path = Path(name)
    parts = {part.lower() for part in path.parts}
    if "__macosx" in parts or ".git" in parts or "attachments" in parts:
        return True
    return any(part.startswith(".") for part in path.parts)


def source_title_from_path(name: str) -> str:
    stem = Path(name).stem.strip() or Path(name).name
    return stem[:255]


def html_to_text(html: str) -> str:
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(html, "html.parser")
        for node in soup(["script", "style", "noscript"]):
            node.decompose()
        return soup.get_text("\n")
    except Exception:
        return html


def detect_language(text: str) -> str:
    lowered = text.lower()
    vietnamese_markers = (
        "\u0103\u00e2\u0111\u00ea\u00f4\u01a1\u01b0"
        "\u00e1\u00e0\u1ea3\u00e3\u1ea1\u1ea5\u1ea7\u1ea9\u1eab\u1ead"
        "\u1eaf\u1eb1\u1eb3\u1eb5\u1eb7\u00e9\u00e8\u1ebb\u1ebd\u1eb9"
        "\u1ebf\u1ec1\u1ec3\u1ec5\u1ec7\u00ed\u00ec\u1ec9\u0129\u1ecb"
        "\u00f3\u00f2\u1ecf\u00f5\u1ecd\u1ed1\u1ed3\u1ed5\u1ed7\u1ed9"
        "\u1edb\u1edd\u1edf\u1ee1\u1ee3\u00fa\u00f9\u1ee7\u0169\u1ee5"
        "\u1ee9\u1eeb\u1eed\u1eef\u1ef1\u00fd\u1ef3\u1ef7\u1ef9\u1ef5"
    )
    if any(char in lowered for char in vietnamese_markers):
        return "vi"
    return "en"


def repair_mojibake(text: str) -> str:
    if not text or not looks_like_mojibake(text):
        return text
    try:
        repaired = text.encode("latin-1").decode("utf-8")
    except UnicodeError:
        return text
    return repaired if mojibake_score(repaired) < mojibake_score(text) else text


def looks_like_mojibake(text: str) -> bool:
    return mojibake_score(text) >= 3


def mojibake_score(text: str) -> int:
    markers = (
        "\u00c3",
        "\u00c2",
        "\u00e2",
        "\u00c4",
        "\u00c6",
        "\u00e1\u00ba",
        "\u00e1\u00bb",
        "\u0080",
        "\u0081",
        "\u0082",
        "\u0083",
    )
    return sum(text.count(marker) for marker in markers)
